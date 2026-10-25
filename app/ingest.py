"""Fully on-device file extraction; OCR via installed local Tesseract only."""
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import os
import shutil
import subprocess

import fitz
from docx import Document
from pptx import Presentation
from PIL import Image, ImageOps

SUPPORTED = {'.pdf', '.docx', '.pptx', '.txt', '.md', '.png', '.jpg', '.jpeg', '.webp'}
MAX_FILE_BYTES = 18 * 1024 * 1024
MAX_PAGES = 100
MAX_OCR_PAGES = 18


@dataclass
class Page:
    number: int
    text: str
    ocr: bool = False


def _ocr_image(raw: bytes) -> str:
    if os.getenv('ENABLE_OCR', 'true').lower() != 'true':
        raise ValueError('Image-only page requires local OCR. Enable OCR in .env.')
    exe = shutil.which('tesseract')
    if not exe:
        raise ValueError('Local OCR is missing. Install Tesseract (brew install tesseract / winget install UB-Mannheim.TesseractOCR).')
    with Image.open(BytesIO(raw)) as image:
        img = ImageOps.exif_transpose(image).convert('RGB')
        img.thumbnail((2100, 2700))
        io = BytesIO()
        img.save(io, format='PNG')
    proc = subprocess.run([exe, 'stdin', 'stdout', '-l', 'eng', '--psm', '3'],
                          input=io.getvalue(), capture_output=True, timeout=90, check=False)
    if proc.returncode:
        raise ValueError('Local OCR could not read the image. Try a clearer image or searchable PDF.')
    text = proc.stdout.decode('utf-8', errors='replace').strip()
    if len(text) < 8:
        raise ValueError('Local OCR found too little readable text. Handwriting and formulas may need manual transcription.')
    return text


def extract_pages(path: Path) -> list[Page]:
    suffix = path.suffix.lower()
    pages = []
    if suffix == '.pdf':
        with fitz.open(path) as pdf:
            if len(pdf) > MAX_PAGES:
                raise ValueError(f'Too many PDF pages (max {MAX_PAGES}).')
            ocr_count = 0
            for i, page in enumerate(pdf):
                text = page.get_text('text', sort=True).strip()
                used_ocr = False
                if len(text) < 35:
                    ocr_count += 1
                    if ocr_count > MAX_OCR_PAGES:
                        raise ValueError(f'Too many scanned PDF pages (max {MAX_OCR_PAGES}).')
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
                    text = _ocr_image(pix.tobytes('png'))
                    used_ocr = True
                if text:
                    pages.append(Page(i+1, text, used_ocr))
    elif suffix == '.docx':
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(' | '.join(c.text for c in row.cells))
        pages.append(Page(1, '\n'.join(lines)))
    elif suffix == '.pptx':
        pres = Presentation(path)
        if len(pres.slides) > MAX_PAGES:
            raise ValueError('Too many slides.')
        for i, slide in enumerate(pres.slides):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    lines.append(shape.text)
                if shape.has_table:
                    lines.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
            if lines:
                pages.append(Page(i+1, '\n'.join(lines)))
    elif suffix in {'.txt', '.md'}:
        content = path.read_text(encoding='utf-8-sig', errors='replace').strip()
        pages = [Page(i // 6500+1, content[i:i+6500]) for i in range(0, len(content), 6500)]
    elif suffix in {'.png', '.jpg', '.jpeg', '.webp'}:
        pages.append(Page(1, _ocr_image(path.read_bytes()), True))
    else:
        raise ValueError('Unsupported file format.')
    if not any(x.text.strip() for x in pages):
        raise ValueError('No readable text. Try a clear searchable file.')
    return pages
